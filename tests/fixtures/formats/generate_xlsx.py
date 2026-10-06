"""Writes the Excel workbook fixtures used by tests/formats-xlsx.test.ts.

openpyxl writes inline strings; XlsxWriter writes shared strings as Excel
does. Run from the repository root with openpyxl and XlsxWriter installed:
    python tests/fixtures/formats/generate_xlsx.py
"""

import math
from datetime import datetime, timedelta
from pathlib import Path

import xlsxwriter
from openpyxl import Workbook

HERE = Path(__file__).parent


def basic():
    book = Workbook()
    run = book.active
    run.title = 'Run 1'
    # Column D has no header; column F stays empty and is not a channel.
    run.append(['Time [s]', 'Torque [Nm]', 'Speed', None, 'Valid'])
    for i in range(300):
        t = i * 0.01
        row = [t, 10 + math.sin(t), i * 2.0, -float(i), i % 2 == 0]
        if i == 5:
            row[1] = None  # empty cell
        if i == 7:
            row[2] = 'n/a'  # text cell
        if i == 9:
            row[1] = '#N/A'  # error cell
        run.append(row)
    run.append([None, 1.0, 2.0])  # no time: skipped
    run.append([3.5, 1.0, 2.0])

    notes = book.create_sheet('Notes')
    notes['A1'] = 'Operator'
    notes['A2'] = 'Kyle'
    book.create_sheet('Blank')

    dates = book.create_sheet('Logged & dated')
    dates.append(['Timestamp', 'Pressure [bar]'])
    start = datetime(2024, 3, 1, 12, 0, 0)
    for i in range(50):
        dates.append([start + timedelta(seconds=0.5 * i), 1 + i / 100])
        dates.cell(row=i + 2, column=1).number_format = 'yyyy-mm-dd hh:mm:ss.000'
    book.save(HERE / 'xlsx-basic.xlsx')


def shared():
    book = xlsxwriter.Workbook(HERE / 'xlsx-shared.xlsx')
    sheet = book.add_worksheet('Rig <A>')
    clock = book.add_format({'num_format': 'h:mm:ss.000'})
    for column, header in enumerate(['Clock', 'Flow [l/min]', 'Flow [l/min]', 'Note']):
        sheet.write_string(0, column, header)
    for i in range(40):
        sheet.write_number(i + 1, 0, (8 * 3600 + i * 0.25) / 86400, clock)
        sheet.write_number(i + 1, 1, i * 1.5)
        # Sparse: column C only every fourth row.
        if i % 4 == 0:
            sheet.write_number(i + 1, 2, -i)
        sheet.write_string(i + 1, 3, 'ok' if i % 2 else 'check')
    book.close()


if __name__ == '__main__':
    basic()
    shared()
